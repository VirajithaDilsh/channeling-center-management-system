#!/usr/bin/env python3
"""Generates testing/Manual_Test_Cases.xlsx from the TEST_CASES data below.
Re-run this script after editing test cases here to regenerate the sheet."""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "Manual_Test_Cases.xlsx")

COLUMNS = [
    "Test Case ID", "Module", "Test Scenario", "Pre-conditions",
    "Test Steps", "Test Data", "Expected Result", "Actual Result",
    "Status", "Priority", "Tested By", "Date Tested", "Remarks",
]

# (module, scenario, preconditions, steps, test_data, expected_result, priority)
TEST_CASES = [
    # ---------------- Login / Authentication ----------------
    ("Login", "Login with valid credentials", "User account exists in the system",
     "1. Go to login page\n2. Enter valid email\n3. Enter valid password\n4. Click Login",
     "Valid email + correct password", "User is authenticated and redirected to the dashboard matching their role", "High"),
    ("Login", "Login with invalid password", "User account exists",
     "1. Go to login page\n2. Enter valid email\n3. Enter wrong password\n4. Click Login",
     "Valid email, wrong password", "Login is rejected with an error message; user stays on login page", "High"),
    ("Login", "Login with unregistered email", "None",
     "1. Go to login page\n2. Enter an email not in the system\n3. Enter any password\n4. Click Login",
     "unknown@test.com / any password", "Error message shown, login denied", "Medium"),
    ("Login", "Login with empty fields", "None",
     "1. Go to login page\n2. Leave email and password blank\n3. Click Login",
     "Empty fields", "Form validation blocks submission with required-field messages", "Medium"),
    ("Login", "Login with SQL/Script injection in fields", "None",
     "1. Go to login page\n2. Enter payload like ' OR 1=1-- in email/password\n3. Click Login",
     "' OR 1=1--", "Login rejected, no data leak or error stack shown to user", "High"),
    ("Login", "Session persists after page refresh", "User is logged in",
     "1. Log in successfully\n2. Refresh the browser page",
     "N/A", "User remains logged in and on the same dashboard route", "Medium"),
    ("Login", "Logout functionality", "User is logged in",
     "1. Click user menu / logout option\n2. Confirm logout",
     "N/A", "User session/token is cleared and user is redirected to login page", "High"),
    ("Login", "Access dashboard route directly without login", "User is logged out",
     "1. Log out or clear browser storage\n2. Manually navigate to /dashboard/patients",
     "URL: /dashboard/patients", "User is redirected to login page, protected route not rendered", "High"),
    ("Login", "Forgot password request with registered email", "Account exists",
     "1. Go to Forgot Password page\n2. Enter registered email\n3. Submit",
     "Registered email", "System accepts request and shows confirmation message", "Medium"),
    ("Login", "Forgot password request with unregistered email", "None",
     "1. Go to Forgot Password page\n2. Enter an email not in system\n3. Submit",
     "unknown@test.com", "Appropriate error/neutral message shown (no account enumeration)", "Low"),

    # ---------------- Admin & Role Management ----------------
    ("Admin Management", "View list of admin/staff users", "Logged in as admin",
     "1. Navigate to Admin page\n2. Observe user list table",
     "N/A", "All registered staff users are listed with name, email, role, contact", "High"),
    ("Admin Management", "Filter/search admin user list", "At least 2+ users exist",
     "1. Go to Admin page\n2. Enter a name/email in search/filter box",
     "Partial name e.g. 'joh'", "List filters to matching users only; clearing filter restores full list", "Medium"),
    ("Admin Management", "Register a new staff user (Register Role)", "Logged in as admin",
     "1. Go to Register Role page\n2. Fill name, email, contact, password, role\n3. Submit",
     "name: Test User, role: billing", "New user created, appears in Admin list, redirected correctly after add", "High"),
    ("Admin Management", "Register new user with duplicate email", "A user with that email already exists",
     "1. Go to Register Role\n2. Enter an already-registered email\n3. Submit",
     "Existing email", "Form shows duplicate-email error; user is not created twice", "High"),
    ("Admin Management", "Register new user with missing required fields", "None",
     "1. Go to Register Role\n2. Leave required fields blank\n3. Submit",
     "Blank fields", "Validation errors shown; submission blocked", "Medium"),
    ("Admin Management", "Edit existing admin/staff details", "At least one user exists",
     "1. Open Edit Admin for a user\n2. Change name/contact/role\n3. Save",
     "Updated contact number", "Changes are saved and reflected in Admin list", "High"),
    ("Admin Management", "Delete an admin/staff user", "At least one non-self user exists",
     "1. Go to Admin page\n2. Click delete on a user\n3. Confirm deletion",
     "N/A", "User is removed from the list and can no longer log in", "High"),
    ("Admin Management", "Role Management - view roles and permissions", "Logged in as admin",
     "1. Navigate to Role Management page\n2. Review roles and assigned permissions",
     "N/A", "Roles (admin, doctor, patient_manager, billing) and their permissions display correctly", "Medium"),
    ("Admin Management", "Role Management - update permissions for a role", "Logged in as admin",
     "1. Open a role\n2. Toggle a permission (e.g. manage_pharmacy)\n3. Save",
     "Toggle manage_pharmacy off", "Permission change is saved; user with that role loses/gains access accordingly", "High"),
    ("Admin Management", "Non-admin cannot access Admin page", "Logged in as non-admin role (e.g. billing)",
     "1. Log in as billing/doctor/patient_manager\n2. Try to navigate to /dashboard/admin",
     "URL: /dashboard/admin", "Access is blocked/redirected; page/menu item not shown", "High"),

    # ---------------- Doctor Management ----------------
    ("Doctor Management", "View list of doctors", "At least one doctor exists",
     "1. Navigate to Doctors page",
     "N/A", "List of doctors displays with name, specialization, contact details", "High"),
    ("Doctor Management", "Add a new doctor with valid details", "Logged in with permission",
     "1. Go to Add Doctor page\n2. Fill in name, specialization, contact, fee, schedule\n3. Submit",
     "name: Dr. Test, fee: 1500", "Doctor is created and appears in Doctor Management list", "High"),
    ("Doctor Management", "Add doctor with missing required fields", "None",
     "1. Go to Add Doctor page\n2. Leave required fields blank\n3. Submit",
     "Blank fields", "Validation prevents submission with clear error messages", "Medium"),
    ("Doctor Management", "Add doctor with invalid contact number format", "None",
     "1. Go to Add Doctor page\n2. Enter letters in contact field\n3. Submit",
     "contact: abc123", "Field validation rejects non-numeric/invalid format", "Low"),
    ("Doctor Management", "View single doctor profile", "Doctor record exists",
     "1. From Doctors list, click on a doctor",
     "N/A", "Doctor View page shows full profile details correctly", "Medium"),
    ("Doctor Management", "Edit doctor details", "Doctor record exists",
     "1. Open Doctor Edit page\n2. Update fee/specialization\n3. Save",
     "Updated fee: 2000", "Changes persist and reflect in Doctor list/profile", "High"),
    ("Doctor Management", "Delete a doctor", "Doctor record exists with no active appointments",
     "1. Go to Doctor Management\n2. Delete a doctor\n3. Confirm",
     "N/A", "Doctor is removed from list; not selectable for new appointments", "Medium"),
    ("Doctor Management", "Doctor Portal - login as doctor and view assigned appointments", "Logged in as doctor role",
     "1. Log in as a doctor\n2. Navigate to Doctor Home",
     "N/A", "Doctor sees only their own appointments for the day/queue", "High"),
    ("Doctor Management", "Doctor Portal - open consultation for an appointment", "An appointment assigned to the doctor exists",
     "1. From Doctor Home, click Consultation on an appointment",
     "N/A", "Consultation page loads with patient details and prescription form", "High"),
    ("Doctor Management", "Doctor Portal - add prescription during consultation", "Consultation page open, medicines exist in inventory",
     "1. In Consultation, search medicine via autocomplete\n2. Select medicine, set quantity/dosage\n3. Save prescription",
     "Medicine: Paracetamol, qty: 10", "Prescription is saved with QUEUED status and appears in Pharmacy Dispense Queue", "High"),
    ("Doctor Management", "Non-doctor cannot access Doctor Portal", "Logged in as non-doctor role",
     "1. Log in as admin/billing/patient_manager\n2. Try to navigate to /dashboard/doctor-home",
     "URL: /dashboard/doctor-home", "Access blocked by ProtectedRoute (requiredPermission doctor_portal)", "High"),

    # ---------------- Patient Management ----------------
    ("Patient Management", "View list of registered patients", "At least one patient exists",
     "1. Navigate to Patients page",
     "N/A", "Patient list displays with name, contact, registration details", "High"),
    ("Patient Management", "Search/filter patients by name or ID", "Multiple patients exist",
     "1. Go to Patients page\n2. Enter search term",
     "Partial patient name", "List filters correctly to matching patients", "Medium"),
    ("Patient Management", "Register a new patient with valid details", "Logged in with manage_patients permission",
     "1. Go to Register Patient page\n2. Fill name, age, gender, contact, address\n3. Submit",
     "name: Test Patient, age: 30", "Patient is created and appears in Patients list", "High"),
    ("Patient Management", "Register patient with missing required fields", "None",
     "1. Go to Register Patient\n2. Leave required fields blank\n3. Submit",
     "Blank fields", "Form validation blocks submission, shows required errors", "Medium"),
    ("Patient Management", "Register patient with invalid contact/NIC format", "None",
     "1. Go to Register Patient\n2. Enter invalid NIC/phone format\n3. Submit",
     "NIC: 123, phone: abcd", "Validation error shown for invalid format", "Low"),
    ("Patient Management", "View patient profile", "Patient record exists",
     "1. From Patients list, click View on a patient",
     "N/A", "View Patient page shows full record: history, appointments, billing summary", "High"),
    ("Patient Management", "Edit patient details", "Patient record exists",
     "1. Open Edit Patient page\n2. Update contact/address\n3. Save",
     "Updated phone number", "Changes are saved and reflected in patient profile", "High"),
    ("Patient Management", "Delete/deactivate a patient record", "Patient record exists",
     "1. From Patients list, delete a patient\n2. Confirm",
     "N/A", "Patient is removed/deactivated; historical records remain intact if applicable", "Medium"),
    ("Patient Management", "Non-permitted role cannot access Patients module", "Logged in as role without manage_patients",
     "1. Log in as a restricted role\n2. Try to navigate to /dashboard/patients",
     "URL: /dashboard/patients", "Access blocked/redirected by ProtectedRoute", "High"),

    # ---------------- Appointments ----------------
    ("Appointments", "View list of appointments", "At least one appointment exists",
     "1. Navigate to Appointments page",
     "N/A", "Appointments list shows patient, doctor, date/time, status", "High"),
    ("Appointments", "Book a new appointment", "Patient and doctor records exist",
     "1. Go to Add Appointment page\n2. Select patient, doctor, date/time\n3. Submit",
     "Existing patient + doctor, future date/time", "Appointment is created; VisitSession auto-opened with Doctor Fee + Center Fee lines", "High"),
    ("Appointments", "Book appointment with a past date/time", "None",
     "1. Go to Add Appointment\n2. Select a date/time in the past\n3. Submit",
     "Past date", "Validation prevents booking a past-dated appointment", "Medium"),
    ("Appointments", "Book overlapping/duplicate appointment for same doctor/slot", "An appointment already exists for that doctor/slot",
     "1. Attempt to book same doctor at an already-booked time slot\n2. Submit",
     "Same doctor, same time slot", "System warns or prevents double-booking for that slot", "Medium"),
    ("Appointments", "Cancel an appointment", "An OPEN appointment exists",
     "1. From Appointments list, select an appointment\n2. Cancel it",
     "N/A", "Appointment status changes to canceled; linked VisitSession reflects CANCELED where applicable", "High"),
    ("Appointments", "Reschedule an appointment", "An existing appointment exists",
     "1. Open an appointment\n2. Change date/time\n3. Save",
     "New valid date/time", "Appointment updates with new schedule, no duplicate created", "Medium"),
    ("Appointments", "Filter appointments by date or doctor", "Multiple appointments exist across dates/doctors",
     "1. Go to Appointments page\n2. Apply date/doctor filter",
     "Specific date or doctor name", "List filters correctly to matching appointments", "Low"),

    # ---------------- Inventory / Pharmacy ----------------
    ("Inventory", "View medicine inventory list", "At least one medicine exists",
     "1. Navigate to Inventory page",
     "N/A", "Medicine list shows name, stock quantity, price, expiry", "High"),
    ("Inventory", "Add a new medicine with valid details", "Logged in with staff role",
     "1. Go to Add Medicine page\n2. Fill name, quantity, price, expiry date\n3. Submit",
     "name: Paracetamol 500mg, qty: 100", "Medicine is added and appears in Inventory list", "High"),
    ("Inventory", "Add medicine with negative or zero stock quantity", "None",
     "1. Go to Add Medicine\n2. Enter negative/zero quantity\n3. Submit",
     "qty: -5", "Validation rejects invalid quantity", "Medium"),
    ("Inventory", "Add medicine with expiry date in the past", "None",
     "1. Go to Add Medicine\n2. Enter a past expiry date\n3. Submit",
     "Expiry: past date", "Validation warns/rejects past expiry date", "Low"),
    ("Inventory", "Edit medicine details / update stock", "Medicine record exists",
     "1. Open Edit Medicine\n2. Update stock quantity or price\n3. Save",
     "Updated qty: 150", "Changes persist and reflect in Inventory list", "High"),
    ("Inventory", "Low stock warning indicator", "A medicine's stock is below threshold",
     "1. Set/observe a medicine with low stock\n2. View Inventory page",
     "Stock below reorder level", "Low-stock medicines are visually flagged/highlighted", "Medium"),
    ("Inventory", "Delete a medicine from inventory", "Medicine record exists, not referenced by pending prescriptions",
     "1. From Inventory list, delete a medicine\n2. Confirm",
     "N/A", "Medicine is removed from list and no longer selectable in prescriptions", "Medium"),

    # ---------------- Pharmacy Dispense Queue ----------------
    ("Pharmacy Dispense Queue", "View queued prescriptions", "A prescription with QUEUED status exists",
     "1. Log in with manage_pharmacy permission\n2. Navigate to Pharmacy Dispense Queue",
     "N/A", "Queued prescriptions list shows patient, medicine, quantity, status", "High"),
    ("Pharmacy Dispense Queue", "Dispense a prescription with sufficient stock", "QUEUED prescription exists, stock available",
     "1. Open a queued prescription\n2. Click Dispense\n3. Confirm",
     "Medicine with stock >= prescribed qty", "Prescription status changes to DISPENSED, stock decrements, bill line posted to VisitSession atomically", "High"),
    ("Pharmacy Dispense Queue", "Attempt to dispense with insufficient stock", "QUEUED prescription exists, stock less than prescribed qty",
     "1. Open a queued prescription where stock < required qty\n2. Attempt to dispense",
     "Prescribed qty > available stock", "System blocks full dispense or offers PARTIAL fulfillment with clear message", "High"),
    ("Pharmacy Dispense Queue", "Partially dispense a prescription", "QUEUED prescription exists, partial stock available",
     "1. Open queued prescription\n2. Dispense less than the prescribed quantity",
     "Dispense qty < prescribed qty", "Prescription status becomes PARTIAL; correct stock decrement and bill line for dispensed amount only", "Medium"),
    ("Pharmacy Dispense Queue", "Reject a prescription", "QUEUED prescription exists",
     "1. Open queued prescription\n2. Click Reject\n3. Provide reason (if required)\n4. Confirm",
     "N/A", "Prescription status changes to REJECTED; no stock/bill changes occur", "Medium"),
    ("Pharmacy Dispense Queue", "Non-permitted role cannot access Pharmacy page", "Logged in as role without manage_pharmacy",
     "1. Log in as doctor/patient_manager without pharmacy permission\n2. Try /dashboard/pharmacy",
     "URL: /dashboard/pharmacy", "Access blocked/redirected by ProtectedRoute", "High"),

    # ---------------- Billing & Payments ----------------
    ("Billing", "View billing list / visit sessions", "At least one visit session exists",
     "1. Navigate to Billing page",
     "N/A", "List shows patient, visit status (OPEN/PENDING_PHARMACY/READY_FOR_PAYMENT/CLOSED), total amount", "High"),
    ("Billing", "Create invoice for a visit ready for payment", "VisitSession status is READY_FOR_PAYMENT",
     "1. Open Billing for a patient\n2. Click Create Invoice\n3. Review line items (Doctor Fee, Center Fee, Pharmacy)\n4. Confirm",
     "N/A", "Invoice is generated with correct itemized totals matching all posted bill lines", "High"),
    ("Billing", "Attempt to create invoice while pharmacy items are still pending", "VisitSession status is PENDING_PHARMACY",
     "1. Open Billing for a visit still awaiting pharmacy dispense\n2. Attempt to create invoice",
     "N/A", "System prevents invoice creation or clearly flags pending pharmacy items until resolved", "High"),
    ("Billing", "Verify bill line accuracy (Doctor Fee + Center Fee + Pharmacy items)", "Appointment created and prescription dispensed",
     "1. Create appointment (auto Doctor Fee + Center Fee)\n2. Dispense a prescription\n3. Open Billing for that visit",
     "N/A", "All bill lines appear correctly with accurate amounts and no duplication", "High"),
    ("Billing", "Close a visit session after payment", "Invoice created and payment recorded",
     "1. Process payment for an invoice\n2. Verify visit session status",
     "N/A", "VisitSession status transitions to CLOSED; visit no longer editable", "High"),
    ("Payments", "Record a payment against an invoice", "An unpaid invoice exists",
     "1. Navigate to Payments page\n2. Select invoice\n3. Enter payment amount/method\n4. Submit",
     "Full payment amount", "Payment is recorded, invoice marked as paid, reflected in reports", "High"),
    ("Payments", "Record a partial payment", "An unpaid invoice exists",
     "1. Go to Payments\n2. Enter an amount less than total due\n3. Submit",
     "Partial amount", "Payment recorded, remaining balance correctly tracked and displayed", "Medium"),
    ("Payments", "Record payment with amount exceeding invoice total", "An unpaid invoice exists",
     "1. Go to Payments\n2. Enter amount greater than invoice total\n3. Submit",
     "Amount > total due", "System rejects or flags overpayment with a validation message", "Medium"),
    ("Payments", "View payment history for a patient", "Patient has prior payments",
     "1. Navigate to patient's billing/payment history",
     "N/A", "All past payments display with date, amount, method, and reference", "Medium"),

    # ---------------- Reports ----------------
    ("Reports", "Generate daily revenue report", "Billing/payment data exists for the day",
     "1. Navigate to Reports page\n2. Select 'Daily' and today's date\n3. Generate report",
     "Today's date", "Report shows accurate totals matching recorded payments for the day", "High"),
    ("Reports", "Generate report for a custom date range", "Billing data exists within range",
     "1. Go to Reports\n2. Select a custom start/end date\n3. Generate",
     "Start: 2026-08-01, End: 2026-08-16", "Report aggregates correctly across the selected range only", "Medium"),
    ("Reports", "Generate report with no data in selected range", "No transactions in a chosen future range",
     "1. Go to Reports\n2. Select a date range with no activity\n3. Generate",
     "Future date range", "Report shows empty state / zero totals without errors", "Low"),
    ("Reports", "Export report (if supported)", "A report has been generated",
     "1. Generate any report\n2. Click Export/Download",
     "N/A", "Report downloads in expected format (PDF/CSV/Excel) with matching data", "Low"),

    # ---------------- Settings & General/Navigation ----------------
    ("Settings", "View and update profile settings", "Logged in user",
     "1. Navigate to Settings page\n2. Update name/contact/password\n3. Save",
     "Updated display name", "Changes are saved and reflected across the app (e.g. Navbar)", "Medium"),
    ("Settings", "Change password with correct current password", "Logged in user",
     "1. Go to Settings\n2. Enter current password + new password\n3. Save",
     "Valid current password + new password", "Password updated; user can log in with new password", "High"),
    ("Settings", "Change password with incorrect current password", "Logged in user",
     "1. Go to Settings\n2. Enter wrong current password + new password\n3. Save",
     "Wrong current password", "Update rejected with clear error message", "Medium"),
    ("Navigation", "Sidebar shows only permitted menu items per role", "Logged in as various roles",
     "1. Log in as each role (admin, doctor, billing, patient_manager)\n2. Observe Sidebar menu",
     "N/A", "Only menu items the role has permission for are visible", "High"),
    ("Navigation", "Notification bell displays relevant alerts", "System has pending items (e.g. low stock, queued prescriptions)",
     "1. Log in\n2. Click Notification Bell icon",
     "N/A", "Relevant notifications display (e.g. low stock, pending pharmacy items)", "Low"),
    ("Navigation", "Unknown route redirects to login/home", "None",
     "1. Navigate to a non-existent URL (e.g. /dashboard/xyz123)",
     "URL: /dashboard/xyz123", "App redirects to '/' (login) instead of showing a blank/broken page", "Low"),
    ("Navigation", "Responsive layout on smaller screen widths", "None",
     "1. Resize browser to tablet/mobile width\n2. Navigate through dashboard pages",
     "N/A", "Layout adjusts (sidebar collapses/hides) without breaking content", "Low"),
]


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Manual Test Cases"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    wrap_align = Alignment(wrap_text=True, vertical="top", horizontal="left")
    center_align = Alignment(wrap_text=True, vertical="top", horizontal="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append(COLUMNS)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    module_counters = {}
    row_idx = 2
    for module, scenario, precond, steps, data, expected, priority in TEST_CASES:
        prefix = "".join(w[0] for w in module.split()).upper()
        module_counters[prefix] = module_counters.get(prefix, 0) + 1
        tc_id = f"TC-{prefix}-{module_counters[prefix]:03d}"

        row = [tc_id, module, scenario, precond, steps, data, expected, "", "Not Executed", priority, "", "", ""]
        ws.append(row)
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = wrap_align
        row_idx += 1

    last_row = row_idx - 1

    # Column widths
    widths = {
        "A": 14, "B": 22, "C": 34, "D": 26, "E": 40, "F": 24,
        "G": 42, "H": 22, "I": 16, "J": 10, "K": 14, "L": 14, "M": 24,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    # Data validation dropdowns for Status and Priority
    status_dv = DataValidation(
        type="list", formula1='"Pass,Fail,Blocked,Not Executed"', allow_blank=True
    )
    ws.add_data_validation(status_dv)
    status_dv.add(f"I2:I{last_row}")

    priority_dv = DataValidation(
        type="list", formula1='"High,Medium,Low"', allow_blank=True
    )
    ws.add_data_validation(priority_dv)
    priority_dv.add(f"J2:J{last_row}")

    # Conditional-ish manual coloring hint row height
    for r in range(2, last_row + 1):
        ws.row_dimensions[r].height = 60

    wb.save(OUTPUT_PATH)
    print(f"Wrote {last_row - 1} test cases to {OUTPUT_PATH}")


if __name__ == "__main__":
    build_workbook()
